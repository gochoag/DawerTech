import csv
import io
import os
from decimal import Decimal, InvalidOperation
from pathlib import Path

from flask import Blueprint, current_app, g, request, send_file
from sqlalchemy import func, or_, select

from dawertech.modules.auth.infrastructure.security import require_auth
from dawertech.modules.devices.infrastructure.orm_models import DeviceORM
from dawertech.modules.inventory.application.dtos import movement_data, part_data
from dawertech.modules.inventory.application.movements_use_case import register_movement
from dawertech.modules.inventory.application.register_part_use_case import RegisterPartUseCase
from dawertech.modules.inventory.application.search_compatibility_use_case import SearchCompatibilityUseCase
from dawertech.modules.inventory.application.sku_generator import generate_smart_sku
from dawertech.modules.inventory.infrastructure.orm_models import CategoryORM, MovementORM, PartORM, VariantORM
from dawertech.modules.inventory.infrastructure.sqlite_repositories import SqliteSparePartRepository
from dawertech.shared.domain.exceptions import ApiProblem, InventoryRuleError
from dawertech.shared.infrastructure.database import db_session, write_session
from dawertech.shared.infrastructure.local_storage import LocalStorageService
from dawertech.shared.presentation.validation import (
    nonnegative_int,
    optional_text,
    pagination,
    payload,
    positive_int,
    required_text,
)

inventory_bp = Blueprint("inventory", __name__)


def selected_devices(session, data: dict) -> list[DeviceORM]:
    identifiers = data.get("device_ids", [])
    if not isinstance(identifiers, list) or len(identifiers) > 100 or any(isinstance(v, bool) or not isinstance(v, int) for v in identifiers):
        raise ApiProblem("invalid_field", "Revisa device_ids")
    unique_ids = set(identifiers)
    devices = list(session.scalars(select(DeviceORM).where(DeviceORM.id.in_(unique_ids)))) if unique_ids else []
    if len(devices) != len(unique_ids):
        raise ApiProblem("invalid_device", "Hay modelos de celular inexistentes")
    return devices


@inventory_bp.get("/parts")
@require_auth()
def list_parts():
    page, limit = pagination()
    search = request.args.get("q", "").strip()[:100]
    device_id = request.args.get("device_id", type=int)
    include_archived = request.args.get("archived") == "true"
    use_case = SearchCompatibilityUseCase()
    return use_case.execute(
        device_id=device_id,
        search_term=search,
        include_archived=include_archived,
        page=page,
        limit=limit,
    )


@inventory_bp.get("/parts/<int:part_id>")
@require_auth()
def get_part(part_id: int):
    with db_session() as session:
        repo = SqliteSparePartRepository(session)
        part = repo.get_by_id(part_id)
        if not part:
            raise ApiProblem("not_found", "Repuesto no encontrado", 404)
        return part_data(part, repo.stock(part.id))


@inventory_bp.get("/parts/suggest-sku")
@require_auth()
def suggest_sku():
    category = request.args.get("category", "")
    name = request.args.get("name", "")
    with db_session() as session:
        return {"sku": generate_smart_sku(session, category, name)}


@inventory_bp.get("/categories")
@inventory_bp.get("/inventory/categories")
@require_auth()
def list_categories():
    with db_session() as session:
        cats = session.scalars(select(CategoryORM).order_by(CategoryORM.name.asc())).all()
        return [{"id": c.id, "name": c.name} for c in cats]


@inventory_bp.post("/categories")
@inventory_bp.post("/inventory/categories")
@require_auth()
def create_category():
    data = payload()
    name = required_text(data, "name", 80)
    with write_session() as session:
        existing = session.scalars(
            select(CategoryORM).where(func.lower(CategoryORM.name) == name.lower())
        ).first()
        if existing:
            return {"id": existing.id, "name": existing.name}, 200
        category = CategoryORM(name=name)
        session.add(category)
        session.flush()
        return {"id": category.id, "name": category.name}, 201


@inventory_bp.get("/variants")
@inventory_bp.get("/inventory/variants")
@require_auth()
def list_variants():
    with db_session() as session:
        variants = session.scalars(select(VariantORM).order_by(VariantORM.name.asc())).all()
        return [{"id": v.id, "name": v.name} for v in variants]


@inventory_bp.post("/variants")
@inventory_bp.post("/inventory/variants")
@require_auth()
def create_variant():
    data = payload()
    name = required_text(data, "name", 160)
    with write_session() as session:
        existing = session.scalars(
            select(VariantORM).where(func.lower(VariantORM.name) == name.lower())
        ).first()
        if existing:
            return {"id": existing.id, "name": existing.name}, 200
        variant = VariantORM(name=name)
        session.add(variant)
        session.flush()
        return {"id": variant.id, "name": variant.name}, 201


@inventory_bp.post("/parts")
@require_auth()
def create_part():
    data = payload()
    sku = optional_text(data, "sku", 64, default="").upper()
    name = required_text(data, "name", 160)
    category = required_text(data, "category", 80)
    variant = optional_text(data, "variant", 160)
    purchase = nonnegative_int(data, "purchase_cents", 0)
    sale = nonnegative_int(data, "sale_cents", 0)
    minimum = nonnegative_int(data, "min_stock", 0)
    initial = nonnegative_int(data, "initial_stock", 0)
    is_universal = bool(data.get("is_universal", False))
    device_ids = data.get("device_ids", [])
    if not isinstance(device_ids, list) or any(isinstance(v, bool) or not isinstance(v, int) for v in device_ids):
        raise ApiProblem("invalid_field", "Revisa device_ids")

    with write_session() as session:
        response = RegisterPartUseCase.execute(
            session,
            sku=sku,
            name=name,
            category=category,
            variant=variant,
            purchase_cents=purchase,
            sale_cents=sale,
            min_stock=minimum,
            initial_stock=initial,
            is_universal=is_universal,
            device_ids=device_ids,
            user_id=g.current_user_id,
        )
    return response, 201


@inventory_bp.patch("/parts/<int:part_id>")
@require_auth()
def update_part(part_id: int):
    data = payload()
    allowed = {"sku", "name", "variant", "category", "purchase_cents", "sale_cents", "min_stock", "device_ids", "archived", "is_universal"}
    if not data or set(data) - allowed:
        raise ApiProblem("invalid_field", "Campos de repuesto inválidos")

    with write_session() as session:
        part = session.get(PartORM, part_id)
        if not part:
            raise ApiProblem("not_found", "Repuesto no encontrado", 404)
        if "sku" in data:
            sku = required_text(data, "sku", 64).upper()
            if sku != part.sku and session.scalar(select(PartORM).where(PartORM.sku == sku)):
                raise ApiProblem("duplicate_sku", "Este SKU ya existe", 409)
            part.sku = sku
        for key, maximum in (("name", 160), ("category", 80)):
            if key in data:
                setattr(part, key, required_text(data, key, maximum))
        if "variant" in data:
            part.variant = optional_text(data, "variant", 160)
        for key in ("purchase_cents", "sale_cents", "min_stock"):
            if key in data:
                setattr(part, key, nonnegative_int(data, key))
        if "is_universal" in data:
            if not isinstance(data["is_universal"], bool):
                raise ApiProblem("invalid_field", "is_universal debe ser booleano")
            part.is_universal = data["is_universal"]
        if "device_ids" in data:
            part.devices = selected_devices(session, data)
        if "archived" in data:
            if not isinstance(data["archived"], bool):
                raise ApiProblem("invalid_field", "Estado inválido")
            part.archived = data["archived"]

        repo = SqliteSparePartRepository(session)
        response = part_data(part, repo.stock(part.id))
    return response


@inventory_bp.get("/movements")
@require_auth()
def list_movements():
    page, limit = pagination()
    part_id = request.args.get("part_id", type=int)
    with db_session() as session:
        query = select(MovementORM)
        if part_id is not None:
            query = query.where(MovementORM.part_id == part_id)
        total = session.scalar(select(func.count()).select_from(query.subquery())) or 0
        movements = session.scalars(query.order_by(MovementORM.created_at.desc(), MovementORM.id.desc()).offset((page - 1) * limit).limit(limit))
        return {"items": [movement_data(m) for m in movements], "total": total, "page": page, "limit": limit}


@inventory_bp.post("/movements")
@require_auth()
def create_movement():
    data = payload()
    part_id = positive_int(data, "part_id")
    kind = required_text(data, "kind", 16)
    reason = required_text(data, "reason", 255)
    if kind not in {"entry", "exit", "adjustment"}:
        raise ApiProblem("invalid_kind", "Tipo de movimiento inválido")
    quantity = None if kind == "adjustment" else positive_int(data, "quantity")
    target = nonnegative_int(data, "target_stock") if kind == "adjustment" else None
    price = nonnegative_int(data, "unit_price_cents") if "unit_price_cents" in data and kind != "adjustment" else None
    try:
        with write_session() as session:
            movement, stock = register_movement(
                session, part_id=part_id, user_id=g.current_user_id, kind=kind, reason=reason,
                quantity=quantity, target_stock=target, unit_price_cents=price,
            )
            response = {"movement": movement_data(movement), "stock": stock}
    except InventoryRuleError as error:
        raise ApiProblem(error.code, error.message, error.status) from error
    return response, 201


@inventory_bp.get("/dashboard")
@require_auth()
def dashboard():
    with db_session() as session:
        repo = SqliteSparePartRepository(session)
        parts = list(session.scalars(select(PartORM).where(PartORM.archived.is_(False))))
        stocks = repo.stocks_for([p.id for p in parts])
        low = [part_data(p, stocks.get(p.id, 0)) for p in parts if stocks.get(p.id, 0) <= p.min_stock]
        recent = session.scalars(select(MovementORM).order_by(MovementORM.created_at.desc(), MovementORM.id.desc()).limit(8))
        return {
            "part_count": len(parts),
            "unit_count": sum(stocks.values()),
            "low_stock_count": len(low),
            "low_stock": low[:10],
            "recent_movements": [movement_data(m) for m in recent],
        }


@inventory_bp.post("/parts/<int:part_id>/image")
@require_auth()
def upload_image(part_id: int):
    upload = request.files.get("image")
    if not upload:
        raise ApiProblem("invalid_image", "Selecciona una imagen")

    storage = LocalStorageService(current_app.config["MEDIA_DIR"])
    image_bytes = upload.read()
    image_name = storage.save_image(image_bytes)

    try:
        with write_session() as session:
            part = session.get(PartORM, part_id)
            if not part:
                raise ApiProblem("not_found", "Repuesto no encontrado", 404)
            previous = part.image_name
            part.image_name = image_name
        if previous:
            storage.delete_image(previous)
    except Exception:
        storage.delete_image(image_name)
        raise
    return {"has_image": True}


@inventory_bp.get("/parts/<int:part_id>/image")
@require_auth()
def get_image(part_id: int):
    with db_session() as session:
        part = session.get(PartORM, part_id)
        if not part or not part.image_name:
            raise ApiProblem("not_found", "Imagen no encontrada", 404)
        image_name = part.image_name

    storage = LocalStorageService(current_app.config["MEDIA_DIR"])
    image_path = storage.get_image_path(image_name)
    if not image_path:
        raise ApiProblem("not_found", "Imagen no encontrada", 404)
    return send_file(os.fspath(image_path), mimetype="image/webp", max_age=3600)


# CSV Import routes
CSV_COLUMNS = ["sku", "name", "variant", "category", "quantity", "purchase_usd", "sale_usd", "min_stock", "compatible_models"]


def price_cents(value: str) -> int:
    try:
        price = Decimal(str(value))
        if not price.is_finite() or price < 0 or price > Decimal("10000000") or price.as_tuple().exponent < -2:
            raise InvalidOperation
        return int(price * 100)
    except (InvalidOperation, ValueError, TypeError):
        raise ValueError("Precio inválido") from None


def whole_number(value) -> int:
    if type(value) is int:
        result = value
    elif isinstance(value, str) and value.isdecimal():
        result = int(value)
    else:
        raise ValueError("Cantidad o mínimo inválido")
    if result < 0 or result > 1_000_000_000:
        raise ValueError("Cantidad o mínimo inválido")
    return result


def normalize_rows(raw_rows: list[dict], session) -> tuple[list[dict], list[dict]]:
    normalized = []
    errors = []
    seen = set()
    if not isinstance(raw_rows, list) or not raw_rows or len(raw_rows) > 500:
        raise ApiProblem("invalid_import", "El CSV debe tener entre 1 y 500 filas")
    for index, row in enumerate(raw_rows, start=2):
        try:
            if not isinstance(row, dict):
                raise ValueError("Fila inválida")
            sku = str(row.get("sku", "")).strip().upper()
            name = str(row.get("name", "")).strip()
            variant = str(row.get("variant", "")).strip()
            category = str(row.get("category", "")).strip()
            if not sku:
                sku = generate_smart_sku(session, category, name)
            if len(sku) > 64 or not name or len(name) > 160 or len(variant) > 160 or not category or len(category) > 80:
                raise ValueError("SKU, nombre, variante o categoría inválidos")
            if sku in seen or session.scalar(select(PartORM.id).where(PartORM.sku == sku)):
                raise ValueError("SKU duplicado")
            seen.add(sku)
            quantity = whole_number(row.get("quantity", 0))
            minimum = whole_number(row.get("min_stock", 0))
            if "purchase_cents" in row or "sale_cents" in row:
                if type(row.get("purchase_cents")) is not int or type(row.get("sale_cents")) is not int:
                    raise ValueError("Precio inválido")
                purchase = row["purchase_cents"]
                sale = row["sale_cents"]
                if purchase < 0 or sale < 0 or purchase > 1_000_000_000 or sale > 1_000_000_000:
                    raise ValueError("Precio inválido")
                device_ids = row.get("device_ids", [])
                if not isinstance(device_ids, list) or len(device_ids) > 100 or any(type(value) is not int for value in device_ids):
                    raise ValueError("Modelos compatibles inválidos")
                known = set(session.scalars(select(DeviceORM.id).where(DeviceORM.id.in_(device_ids))))
                if known != set(device_ids):
                    raise ValueError("Celular no encontrado")
            else:
                purchase = price_cents(row.get("purchase_usd", "0"))
                sale = price_cents(row.get("sale_usd", "0"))
                compatible = row.get("compatible_models", "")
                if not isinstance(compatible, str):
                    raise ValueError("Modelos compatibles inválidos")
                device_ids = []
                for model in [item.strip() for item in compatible.split(";") if item.strip()]:
                    if "|" not in model:
                        raise ValueError("Usa Marca|Modelo para cada celular compatible")
                    brand, model_name = [part.strip() for part in model.split("|", 1)]
                    device_id = session.scalar(select(DeviceORM.id).where(func.lower(DeviceORM.brand) == brand.lower(), func.lower(DeviceORM.name) == model_name.lower()))
                    if device_id is None:
                        raise ValueError(f"Celular no encontrado: {model}")
                    device_ids.append(device_id)
            normalized.append({"sku": sku, "name": name, "variant": variant, "category": category, "quantity": quantity, "purchase_cents": purchase, "sale_cents": sale, "min_stock": minimum, "device_ids": sorted(set(device_ids))})
        except (ValueError, TypeError) as exc:
            errors.append({"row": index, "message": str(exc)})
    return normalized, errors


@inventory_bp.get("/imports/template")
@require_auth()
def import_template():
    example = ["", "Display OLED", "Negro", "Pantallas", "3", "25.50", "42.00", "1", "Apple|iPhone 13;Apple|iPhone 13 Pro"]
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(CSV_COLUMNS)
    writer.writerow(example)
    output.seek(0)
    return send_file(io.BytesIO(output.getvalue().encode("utf-8")), mimetype="text/csv", as_attachment=True, download_name="plantilla_repuestos.csv")


@inventory_bp.get("/imports/template.xlsx")
@require_auth()
def import_excel_template():
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Catálogo Repuestos"

    labels = [
        "sku",
        "name",
        "variant",
        "category",
        "quantity",
        "purchase_usd",
        "sale_usd",
        "min_stock",
        "compatible_models",
    ]
    display_headers = [
        "SKU (Opcional)",
        "Nombre Repuesto *",
        "Variante / Calidad",
        "Categoría *",
        "Cantidad Inicial",
        "Precio Compra USD",
        "Precio Venta USD",
        "Stock Mínimo",
        "Modelos Compatibles (Marca|Modelo separados por ;)",
    ]

    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2A2F63", end_color="2A2F63", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin", color="000000"),
        right=Side(style="thin", color="000000"),
        top=Side(style="thin", color="000000"),
        bottom=Side(style="thin", color="000000"),
    )

    ws.append(labels)
    for col_num in range(1, len(labels) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    ws.row_dimensions[1].height = 26

    sample_rows = [
        ["", "Pantalla OLED Incell", "Negro", "Pantallas", 5, 22.50, 45.00, 2, "Apple|iPhone 13;Apple|iPhone 13 Pro"],
        ["", "Batería Galaxy S21", "Original", "Baterías", 4, 12.00, 25.00, 1, "Samsung|Galaxy S21"],
        ["", "Pin de carga tipo C", "Universal", "Pines de carga", 20, 1.50, 5.00, 5, ""],
    ]

    for r_idx, row in enumerate(sample_rows, start=2):
        ws.append(row)
        for c_idx in range(1, len(row) + 1):
            cell = ws.cell(row=r_idx, column=c_idx)
            cell.font = Font(name="Arial", size=10)
            cell.border = thin_border
            if c_idx in (5, 8):
                cell.alignment = Alignment(horizontal="center")
            elif c_idx in (6, 7):
                cell.number_format = '$#,##0.00'
                cell.alignment = Alignment(horizontal="right")
        ws.row_dimensions[r_idx].height = 20

    widths = [18, 30, 20, 20, 16, 18, 18, 14, 45]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    ws.views.sheetView[0].showGridLines = True

    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return send_file(
        bio,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name="plantilla_repuestos_dawertech.xlsx",
    )


@inventory_bp.post("/imports/preview")
@require_auth()
def preview_import():
    file = request.files.get("file")
    if not file:
        raise ApiProblem("invalid_file", "Adjunta un archivo Excel (.xlsx) o CSV (.csv)")

    filename = (file.filename or "").lower()
    raw_rows = []

    is_excel = (
        filename.endswith(".xlsx")
        or filename.endswith(".xls")
        or (file.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" and not filename.endswith(".csv"))
    )

    if is_excel:
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(file.read()), data_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                raise ApiProblem("invalid_file", "El archivo Excel está vacío")
            header_row = [str(c).strip().lower() if c is not None else "" for c in rows[0]]
            col_map = {}
            for col_idx, h in enumerate(header_row):
                for k in CSV_COLUMNS:
                    if k in h:
                        col_map[k] = col_idx
                        break
            if len(col_map) < 4:
                col_map = {col: i for i, col in enumerate(CSV_COLUMNS) if i < len(header_row)}
            for r in rows[1:]:
                if not any(r):
                    continue
                row_dict = {}
                for k in CSV_COLUMNS:
                    val = r[col_map[k]] if k in col_map and col_map[k] < len(r) else ""
                    row_dict[k] = str(val).strip() if val is not None else ""
                raw_rows.append(row_dict)
        except Exception as exc:
            if isinstance(exc, ApiProblem):
                raise
            raise ApiProblem("invalid_file", f"Error al leer el archivo Excel: {str(exc)}") from exc
    else:
        try:
            content = file.read().decode("utf-8-sig")
        except UnicodeDecodeError:
            raise ApiProblem("invalid_encoding", "El archivo debe estar en codificación UTF-8") from None
        reader = csv.DictReader(io.StringIO(content))
        if not reader.fieldnames or [col.strip().lower() for col in reader.fieldnames] != CSV_COLUMNS:
            raise ApiProblem("invalid_columns", "Las columnas del archivo no coinciden con la plantilla")
        raw_rows = [{k.strip().lower(): (v.strip() if isinstance(v, str) else v) for k, v in row.items()} for row in reader]

    with db_session() as session:
        normalized, errors = normalize_rows(raw_rows, session)
    return {"rows": normalized, "errors": errors, "count": len(normalized)}


@inventory_bp.post("/imports/commit")
@require_auth()
def commit_import():
    data = payload()
    raw_rows = data.get("rows", [])
    with write_session() as session:
        normalized, errors = normalize_rows(raw_rows, session)
        if errors:
            raise ApiProblem("validation_error", "Corrige las filas antes de importar", 422, {"errors": errors})
        parts = []
        for row in normalized:
            devices = list(session.scalars(select(DeviceORM).where(DeviceORM.id.in_(row["device_ids"])))) if row["device_ids"] else []
            part = PartORM(
                sku=row["sku"], name=row["name"], variant=row["variant"], category=row["category"],
                purchase_cents=row["purchase_cents"], sale_cents=row["sale_cents"], min_stock=row["min_stock"],
            )
            part.devices = devices
            session.add(part)
            session.flush()
            if row["quantity"]:
                session.add(MovementORM(
                    part_id=part.id, user_id=g.current_user_id, kind="entry",
                    quantity_delta=row["quantity"], reason="Carga masiva", unit_price_cents=row["purchase_cents"],
                ))
            parts.append(part_data(part, row["quantity"]))
    return {"items": parts, "count": len(parts)}, 201


