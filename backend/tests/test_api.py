import io
import sqlite3
import zipfile
from concurrent.futures import ThreadPoolExecutor

import pytest
from PIL import Image
from sqlalchemy import select

from dawertech import create_app
from dawertech.shared.infrastructure.database import Base, db_session, write_session
from dawertech.modules.auth.infrastructure.orm_models import User
from dawertech.modules.auth.infrastructure.security import hash_password


@pytest.fixture
def app(tmp_path):
    app = create_app({
        "TESTING": True,
        "DATABASE_URL": f"sqlite:///{(tmp_path / 'inventory.sqlite3').as_posix()}",
        "MEDIA_DIR": str(tmp_path / "media"),
        "SECRET_KEY": "test-secret-key-of-sufficient-length",
    })
    Base.metadata.create_all(app.extensions["dawertech_engine"])
    with app.app_context(), write_session() as session:
        session.add(User(email="owner@example.com", name="Dueño", password_hash=hash_password("owner-password-123"), role="owner"))
    yield app
    app.extensions["dawertech_engine"].dispose()


def login(client, email="owner@example.com", password="owner-password-123"):
    response = client.post("/api/v1/auth/login", json={"email": email, "password": password})
    assert response.status_code == 200
    return response.json


def authorized(client, token, path, method="get", **kwargs):
    return getattr(client, method)(path, headers={"Authorization": f"Bearer {token}"}, **kwargs)


def part_payload(**extra):
    return {"sku": "DISP-1", "name": "Display", "category": "Pantallas", "variant": "OLED", "purchase_cents": 1250, "sale_cents": 2300, "min_stock": 1, "initial_stock": 2, "device_ids": [], **extra}


def test_inventory_compatibility_and_negative_stock(app):
    client = app.test_client()
    token = login(client)["access_token"]
    device = authorized(client, token, "/api/v1/devices", method="post", json={"brand": "Apple", "name": "iPhone 13", "code": "A2633"})
    assert device.status_code == 201
    part = authorized(client, token, "/api/v1/parts", method="post", json=part_payload(device_ids=[device.json["id"]]))
    assert part.status_code == 201
    assert part.json["stock"] == 2
    assert part.json["device_ids"] == [device.json["id"]]
    matched = authorized(client, token, f"/api/v1/parts?device_id={device.json['id']}")
    assert matched.json["total"] == 1
    insufficient = authorized(client, token, "/api/v1/movements", method="post", json={"part_id": part.json["id"], "kind": "exit", "quantity": 3, "reason": "Reparación"})
    assert insufficient.status_code == 409
    assert authorized(client, token, f"/api/v1/parts/{part.json['id']}").json["stock"] == 2
    adjustment = authorized(client, token, "/api/v1/movements", method="post", json={"part_id": part.json["id"], "kind": "adjustment", "target_stock": 1, "reason": "Conteo"})
    assert adjustment.status_code == 201
    assert adjustment.json["movement"]["quantity_delta"] == -1
    assert authorized(client, token, f"/api/v1/parts/{part.json['id']}").json["stock"] == 1


def test_user_limit_roles_password_change_and_revocation(app):
    client = app.test_client()
    owner = login(client)
    token = owner["access_token"]
    created = authorized(client, token, "/api/v1/users", method="post", json={"name": "Ana", "email": "ana@example.com", "password": "initial-password-123"})
    assert created.status_code == 201
    coadmin = login(client, "ana@example.com", "initial-password-123")
    co_token = coadmin["access_token"]
    assert authorized(client, co_token, "/api/v1/parts").status_code == 403
    changed = authorized(client, co_token, "/api/v1/auth/password", method="post", json={"old_password": "initial-password-123", "new_password": "another-password-123"})
    assert changed.status_code == 204
    assert authorized(client, co_token, "/api/v1/parts").status_code == 401
    coadmin = login(client, "ana@example.com", "another-password-123")
    co_token = coadmin["access_token"]
    assert authorized(client, co_token, "/api/v1/parts").status_code == 200
    assert authorized(client, co_token, "/api/v1/users").status_code == 403
    for index in range(3):
        response = authorized(client, token, "/api/v1/users", method="post", json={"name": f"User {index}", "email": f"user{index}@example.com", "password": "initial-password-123"})
        assert response.status_code == 201
    assert authorized(client, token, "/api/v1/users", method="post", json={"name": "Sixth", "email": "sixth@example.com", "password": "initial-password-123"}).status_code == 409
    assert authorized(client, token, f"/api/v1/users/{created.json['id']}", method="patch", json={"active": False}).status_code == 200
    assert authorized(client, co_token, "/api/v1/parts").status_code == 401
    assert client.post("/api/v1/auth/refresh", json={"refresh_token": coadmin["refresh_token"]}).status_code == 401


def test_csv_preview_commit_and_atomic_validation(app):
    client = app.test_client()
    token = login(client)["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    device = authorized(client, token, "/api/v1/devices", method="post", json={"brand": "Apple", "name": "iPhone 13", "code": ""}).json
    csv_data = b"sku,name,variant,category,quantity,purchase_usd,sale_usd,min_stock,compatible_models\nA-1,Battery,Original,Batteries,3,2.50,4.00,1,Apple|iPhone 13\n"
    preview = client.post("/api/v1/imports/preview", headers=headers, data={"file": (io.BytesIO(csv_data), "items.csv")})
    assert preview.status_code == 200
    assert preview.json["rows"][0]["purchase_cents"] == 250
    assert preview.json["rows"][0]["device_ids"] == [device["id"]]
    committed = client.post("/api/v1/imports/commit", headers=headers, json={"rows": preview.json["rows"]})
    assert committed.status_code == 201
    assert committed.json["items"][0]["purchase_cents"] == 250
    assert committed.json["items"][0]["stock"] == 3
    assert committed.json["items"][0]["device_ids"] == [device["id"]]
    bad = [{**part_payload(sku="B-1"), "quantity": 1}, {**part_payload(sku="B-2"), "quantity": -1}]
    failed = client.post("/api/v1/imports/commit", headers=headers, json={"rows": bad})
    assert failed.status_code == 422
    assert authorized(client, token, "/api/v1/parts?q=B-1").json["total"] == 0


def test_simultaneous_exits_never_make_stock_negative(app):
    client = app.test_client()
    token = login(client)["access_token"]
    part = authorized(client, token, "/api/v1/parts", method="post", json=part_payload(initial_stock=3)).json
    tokens = [token]
    for index in range(4):
        email = f"colleague{index}@example.com"
        assert authorized(client, token, "/api/v1/users", method="post", json={"name": f"Colega {index}", "email": email, "password": "temporary-password-123"}).status_code == 201
        temporary = login(client, email, "temporary-password-123")["access_token"]
        assert authorized(client, temporary, "/api/v1/auth/password", method="post", json={"old_password": "temporary-password-123", "new_password": "personal-password-123"}).status_code == 204
        tokens.append(login(client, email, "personal-password-123")["access_token"])

    def exit_once(actor_token):
        local_client = app.test_client()
        return authorized(local_client, actor_token, "/api/v1/movements", method="post", json={"part_id": part["id"], "kind": "exit", "quantity": 1, "reason": "Uso"}).status_code

    with ThreadPoolExecutor(max_workers=5) as pool:
        statuses = list(pool.map(exit_once, tokens))
    assert sorted(statuses) == [201, 201, 201, 409, 409]
    assert authorized(client, token, f"/api/v1/parts/{part['id']}").json["stock"] == 0


def test_backup_restores_data(app, tmp_path):
    client = app.test_client()
    token = login(client)["access_token"]
    part = authorized(client, token, "/api/v1/parts", method="post", json=part_payload()).json
    image = io.BytesIO()
    Image.new("RGB", (50, 50), "#5FB0C9").save(image, "PNG")
    image.seek(0)
    uploaded = authorized(client, token, f"/api/v1/parts/{part['id']}/image", method="post", data={"image": (image, "display.png")})
    assert uploaded.status_code == 200
    assert authorized(client, token, f"/api/v1/parts/{part['id']}/image").mimetype == "image/webp"
    runner = app.test_cli_runner()
    destination = tmp_path / "backup"
    result = runner.invoke(args=["respaldar", str(destination)])
    assert result.exit_code == 0, result.output
    backup = next(destination.glob("*.sqlite3"))
    with zipfile.ZipFile(next(destination.glob("*.zip"))) as media:
        assert any(name.endswith(".webp") for name in media.namelist())
    with sqlite3.connect(backup) as connection:
        assert connection.execute("PRAGMA integrity_check").fetchone()[0] == "ok"
        assert connection.execute("SELECT COUNT(*) FROM parts").fetchone()[0] == 1
        restored = tmp_path / "restored.sqlite3"
        with sqlite3.connect(restored) as target:
            connection.backup(target)
            assert target.execute("PRAGMA integrity_check").fetchone()[0] == "ok"
            assert target.execute("SELECT COUNT(*) FROM parts").fetchone()[0] == 1


def test_universal_parts_compatibility_rules(app):
    client = app.test_client()
    token = login(client)["access_token"]

    dev_a = authorized(client, token, "/api/v1/devices", method="post", json={"brand": "Samsung", "name": "Galaxy S22", "code": "SM-S901B"}).json
    dev_b = authorized(client, token, "/api/v1/devices", method="post", json={"brand": "Xiaomi", "name": "Redmi Note 11", "code": "2201117TG"}).json

    # Repuesto exclusivo / compatible solo con Samsung
    part_samsung = authorized(client, token, "/api/v1/parts", method="post", json={
        "sku": "SAM-S22-BATT", "name": "Batería EB-BS901ABY", "category": "Baterías",
        "purchase_cents": 1000, "sale_cents": 2500, "min_stock": 2, "initial_stock": 5,
        "is_universal": False, "device_ids": [dev_a["id"]],
    }).json
    assert part_samsung["is_universal"] is False
    assert part_samsung["compatible_count"] == 1

    # Repuesto universal / insumo general (alcohol, cinta, herramienta)
    part_universal = authorized(client, token, "/api/v1/parts", method="post", json={
        "sku": "INS-ALC-01", "name": "Alcohol Isopropílico 1L", "category": "Químicos",
        "purchase_cents": 300, "sale_cents": 600, "min_stock": 1, "initial_stock": 10,
        "is_universal": True, "device_ids": [],
    }).json
    assert part_universal["is_universal"] is True

    # Búsqueda filtrada por Samsung: debe incluir el repuesto específico Y el insumo universal
    res_a = authorized(client, token, f"/api/v1/parts?device_id={dev_a['id']}").json
    skus_a = {item["sku"] for item in res_a["items"]}
    assert "SAM-S22-BATT" in skus_a
    assert "INS-ALC-01" in skus_a

    # Búsqueda filtrada por Xiaomi: debe incluir el insumo universal pero NO la batería de Samsung
    res_b = authorized(client, token, f"/api/v1/parts?device_id={dev_b['id']}").json
    skus_b = {item["sku"] for item in res_b["items"]}
    assert "INS-ALC-01" in skus_b
    assert "SAM-S22-BATT" not in skus_b


def test_smart_sku_automatic_generation(app):
    client = app.test_client()
    token = login(client)["access_token"]

    # Endpoint de sugerencia previa
    suggest = authorized(client, token, "/api/v1/parts/suggest-sku?category=Pantallas&name=Display iPhone 13 OLED").json
    assert "sku" in suggest
    assert suggest["sku"].startswith("PAN-IP13-")

    # Creación sin enviar SKU: el backend debe autogenerarlo inteligentemente
    part1 = authorized(client, token, "/api/v1/parts", method="post", json={
        "name": "Batería Samsung S21 Ultra",
        "category": "Baterías",
        "purchase_cents": 1500,
        "sale_cents": 2800,
        "min_stock": 2,
    }).json
    assert part1["sku"].startswith("BAT-S21-")
    assert len(part1["sku"]) >= 8

    # Creación de otro repuesto similar: debe tener SKU único sin duplicidad
    part2 = authorized(client, token, "/api/v1/parts", method="post", json={
        "name": "Batería Samsung S21 Plus",
        "category": "Baterías",
        "purchase_cents": 1400,
        "sale_cents": 2600,
        "min_stock": 1,
    }).json
    assert part2["sku"].startswith("BAT-S21-")
    assert part1["sku"] != part2["sku"]


def test_duplicate_device_and_image_upload(app):
    client = app.test_client()
    token = login(client)["access_token"]

    # 1. Registro exitoso de dispositivo
    created = authorized(client, token, "/api/v1/devices", method="post", json={
        "brand": "Samsung",
        "name": "Galaxy S24 Ultra",
        "code": "SM-S928B"
    })
    assert created.status_code == 201
    dev_id = created.json["id"]
    assert created.json["has_image"] is False

    # 2. Intento de registrar duplicado (mismo brand y name)
    duplicate = authorized(client, token, "/api/v1/devices", method="post", json={
        "brand": "Samsung",
        "name": "Galaxy S24 Ultra",
        "code": "SM-S928B-DUP"
    })
    assert duplicate.status_code == 409
    error_msg = duplicate.json["error"]["message"]
    assert "Samsung Galaxy S24 Ultra" in error_msg
    assert "ya se encuentra registrado" in error_msg

    # 3. Subida de imagen al dispositivo
    img = Image.new("RGB", (100, 100), color="blue")
    img_byte_arr = io.BytesIO()
    img.save(img_byte_arr, format="PNG")
    img_byte_arr.seek(0)

    upload_res = client.post(
        f"/api/v1/devices/{dev_id}/image",
        headers={"Authorization": f"Bearer {token}"},
        data={"file": (img_byte_arr, "s24.png")}
    )
    assert upload_res.status_code == 200

    # 4. Obtención de imagen
    get_img = client.get(f"/api/v1/devices/{dev_id}/image", headers={"Authorization": f"Bearer {token}"})
    assert get_img.status_code == 200
    assert get_img.content_type == "image/webp"

    # 5. Consulta de dispositivo refleja has_image = True
    dev_detail = authorized(client, token, f"/api/v1/devices/{dev_id}")
    assert dev_detail.status_code == 200
    assert dev_detail.json["has_image"] is True


def test_excel_template_and_xlsx_preview(app):
    import openpyxl

    client = app.test_client()
    token = login(client)["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    # 1. Descarga de plantilla Excel estilizada
    template_res = client.get("/api/v1/imports/template.xlsx", headers=headers)
    assert template_res.status_code == 200
    assert "spreadsheetml.sheet" in template_res.content_type

    # Verificar que el contenido es un workbook válido
    wb_downloaded = openpyxl.load_workbook(io.BytesIO(template_res.data))
    assert "Catálogo Repuestos" in wb_downloaded.sheetnames

    # 2. Creación de un archivo Excel .xlsx válido en memoria
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Importación"
    ws.append(["sku", "name", "variant", "category", "quantity", "purchase_usd", "sale_usd", "min_stock", "compatible_models"])
    ws.append(["XLS-BAT-01", "Batería Pro Max", "Original", "Baterías", 5, 12.50, 25.00, 2, ""])
    xlsx_stream = io.BytesIO()
    wb.save(xlsx_stream)
    xlsx_stream.seek(0)

    # 3. Vista previa del archivo Excel
    preview_res = client.post("/api/v1/imports/preview", headers=headers, data={
        "file": (xlsx_stream, "catalogo_repuestos.xlsx")
    })
    assert preview_res.status_code == 200
    preview_data = preview_res.json
    assert preview_data["count"] == 1
    assert preview_data["rows"][0]["sku"] == "XLS-BAT-01"
    assert preview_data["rows"][0]["purchase_cents"] == 1250
    assert preview_data["rows"][0]["sale_cents"] == 2500

    # 4. Confirmación de importación
    commit_res = client.post("/api/v1/imports/commit", headers=headers, json={"rows": preview_data["rows"]})
    assert commit_res.status_code == 201
    assert commit_res.json["count"] == 1


def test_categories_and_variants_endpoints(app):
    client = app.test_client()
    tokens = login(client)
    headers = {"Authorization": f"Bearer {tokens['access_token']}"}

    # 1. Categories listing and creation
    res = client.get("/api/v1/categories", headers=headers)
    assert res.status_code == 200
    cats = [c["name"] for c in res.json]
    assert "Pantallas" in cats
    assert "Baterías" in cats

    res_post = client.post("/api/v1/categories", headers=headers, json={"name": "Placas Madre"})
    assert res_post.status_code == 201
    assert res_post.json["name"] == "Placas Madre"
    assert res_post.json["id"] > 0

    # Idempotent re-creation
    res_repeat = client.post("/api/v1/categories", headers=headers, json={"name": "placas madre"})
    assert res_repeat.status_code == 200
    assert res_repeat.json["id"] == res_post.json["id"]

    # 2. Variants listing and creation
    res_var = client.get("/api/v1/variants", headers=headers)
    assert res_var.status_code == 200
    vars = [v["name"] for v in res_var.json]
    assert "Original Service Pack" in vars

    res_var_post = client.post("/api/v1/variants", headers=headers, json={"name": "AAA Grade+"})
    assert res_var_post.status_code == 201
    assert res_var_post.json["name"] == "AAA Grade+"

    res_var_repeat = client.post("/api/v1/variants", headers=headers, json={"name": "aaa grade+"})
    assert res_var_repeat.status_code == 200
    assert res_var_repeat.json["id"] == res_var_post.json["id"]



